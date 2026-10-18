from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.config import ROOT
from app.models import Evidence

COMMERCIAL = ["product_id", "Brand", "Manufacturer Part #", "Ferguson Part #", "UPC",
              "Price (USD)", "Availability", "Ferguson URL", "Spec Sheet URL"]


def safe_cell(cell, value, formula=False):
    cell.value = value
    # External text is always text, even when it begins with =, +, -, or @.
    if isinstance(value, str) and not formula:
        cell.data_type = "s"


def picture_formula(url):
    if not url.startswith("https://") or any(x in url for x in ('"', '\n', '\r')):
        return None
    return f'=IMAGE("{url}")'


REVIEW_COLUMNS = ["#", "product_id", "field", "verdict", "severity", "finding", "evidence", "current_value",
                  "proposed_value", "proposed_source", "found_by", "decision"]
SEVERITY_ORDER = {"critical": 0, "warning": 1, "info": 2}


def write_workbook(products, request, path, mode, review="", findings=None, agent_summary=None):
    workbook = load_workbook(ROOT / "templates/rfq_template.xlsx")
    rfq, sources = workbook["RFQ"], workbook["Sources"]
    headers = [cell.value for cell in rfq[1]]
    notes = workbook.create_sheet("Notes")
    notes.append(["product_id", "note"])
    summary = workbook.create_sheet("Summary")
    summary.append(["Metric", "Value"])
    agent_sheet = workbook.create_sheet("Agent Review")
    agent_sheet.append(REVIEW_COLUMNS)
    commercial = workbook.create_sheet("Commercial")
    commercial.append(COMMERCIAL)
    # Citations for RFQ keep the template's column naming. Extra sheet fields are qualified.
    source_rows = []
    populated, possible = 0, len(products) * (len(headers) - 2)
    for row_index, product in enumerate(products, 2):
        identity = Evidence(product.product_id, "email", request.origins[product.product_id])
        status = Evidence("Found" if product.found else "Not Found", "workflow",
                          "app/lookup.py; " + ("; ".join(product.notes) or "exact live product match"))
        fields = dict(product.fields)
        fields["product_id"], fields["Lookup Status"] = identity, status
        if "Picture URL" in fields:
            fields["Picture"] = fields["Picture URL"]
            formula = picture_formula(fields["Picture"].value)
            if formula:
                evidence = fields["Picture"]
                fields["Image"] = Evidence(formula, evidence.source_type, evidence.source, evidence.confidence)
        missing = []
        for column_index, name in enumerate(headers, 1):
            evidence = fields.get(name)
            if evidence:
                cell = rfq.cell(row_index, column_index)
                safe_cell(cell, evidence.value, formula=name == "Image")
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                source_rows.append((product.product_id, name, evidence.value, evidence.source_type,
                                    evidence.confidence, evidence.source))
                if name not in ("product_id", "Lookup Status"):
                    populated += 1
            elif name not in ("product_id", "Lookup Status"):
                missing.append(name)
        if missing:
            note = "Missing values (left blank): " + ", ".join(missing)
            if note not in product.notes:
                product.notes.append(note)
        for col, name in enumerate(COMMERCIAL, 1):
            evidence = fields.get(name)
            if evidence:
                safe_cell(commercial.cell(row_index, col), evidence.value)
                source_rows.append((product.product_id, "Commercial." + name, evidence.value,
                                    evidence.source_type, evidence.confidence, evidence.source))
    for row in source_rows:
        target = sources.max_row + 1
        for col, value in enumerate(row, 1):
            safe_cell(sources.cell(target, col), value)
    findings = findings or []
    # Findings keep their index (the id the dashboard approves by); sorted most severe first.
    ordered = sorted(enumerate(findings, 1), key=lambda item: (SEVERITY_ORDER.get(item[1].severity, 3),
                                                               item[1].verdict == "inferred", item[0]))
    for number, finding in ordered:
        row = agent_sheet.max_row + 1
        values = [number, finding.product_id, finding.field, finding.verdict, finding.severity, finding.message,
                  finding.evidence, finding.current_value, finding.proposed_value, finding.proposed_source,
                  finding.found_by, finding.decision or ("pending" if finding.proposed_value else "")]
        for col, value in enumerate(values, 1):
            safe_cell(agent_sheet.cell(row, col), value)
    agent_notes = [(f.product_id, f"Agent review ({f.severity}): {f.field}: {f.message}"
                    + (f" Proposed: {f.proposed_value!r} ({f.decision or 'awaiting approval'})." if f.proposed_value else ""))
                   for f in findings if f.severity in ("critical", "warning")]
    all_notes = request.notes + [(p.product_id, n) for p in products for n in p.notes] + agent_notes
    for product_id, note in all_notes:
        row = notes.max_row + 1
        safe_cell(notes.cell(row, 1), product_id)
        safe_cell(notes.cell(row, 2), note)
    metrics = {
        "Products Processed": len(products),
        "Products Found": sum(p.found for p in products),
        "Not Found": sum(not p.found for p in products),
        "Fields Completed": f"{populated / possible:.1%}" if possible else "0.0%",
        "Warnings / Notes": len(all_notes),
        "Quality Mode": mode,
        "Coverage Definition": "Populated RFQ specification cells / all requested specification cells; includes unknown products.",
    }
    if agent_summary:
        reviewed = agent_summary["values_reviewed"] or 1
        metrics.update({
            "Agent: values verified": f"{agent_summary['verified']} of {agent_summary['values_reviewed']} "
                                      f"({agent_summary['verified'] / reviewed:.0%})",
            "Agent: values inferred": agent_summary["inferred"],
            "Agent: source conflicts": agent_summary["conflicts"],
            "Agent: issues (critical / warning)": f"{agent_summary['issues']} "
                                                  f"({agent_summary['critical']} / {agent_summary['warnings']})",
            "Agent: fixes proposed": agent_summary["fixes_proposed"],
            "Agent: fixes approved": sum(1 for f in findings if f.decision == "approved"),
        })
    if review:
        metrics["AI Review (advisory)"] = review
    for key, value in metrics.items():
        row = summary.max_row + 1
        safe_cell(summary.cell(row, 1), key)
        safe_cell(summary.cell(row, 2), value)
    for sheet in workbook:
        if sheet.title == "How to fill":
            continue
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="17365D")
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 24
    notes.column_dimensions["B"].width = 110
    for letter, width in zip("ABCDEFGHIJKL", (5, 14, 18, 11, 10, 70, 60, 26, 22, 50, 10, 11)):
        agent_sheet.column_dimensions[letter].width = width
    sources.column_dimensions["F"].width = 100
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 100
    path = Path(path)
    workbook.save(path)
    workbook.close()
    validate_workbook(path, request.ids)
    return metrics


def validate_workbook(path, expected_ids):
    """Reopen the actual XLSX and reject missing rows or missing citations."""
    workbook = load_workbook(path)
    try:
        actual = [row[0] for row in list(workbook["RFQ"].values)[1:]]
        if actual != expected_ids:
            raise ValueError("Workbook product order/coverage does not match the request")
        citations = {(r[0], r[1], r[2]) for r in list(workbook["Sources"].values)[1:] if r[5]}
        for title in ("RFQ", "Commercial"):
            rows = list(workbook[title].values)
            for row in rows[1:]:
                for name, value in zip(rows[0], row):
                    field = name if title == "RFQ" else f"Commercial.{name}"
                    if value is not None and (row[0], field, value) not in citations:
                        raise ValueError(f"Missing source for {row[0]} / {field}")
    finally:
        workbook.close()
