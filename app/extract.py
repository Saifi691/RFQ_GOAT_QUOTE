"""Read MIME bodies and XLSX item columns without executing attachment content."""
import re
from dataclasses import dataclass
from email import policy
from email.parser import BytesParser
from io import BytesIO
from zipfile import ZipFile, BadZipFile

from bs4 import BeautifulSoup
from openpyxl import load_workbook

MAX_EMAIL_BYTES = 15 * 1024 * 1024
MAX_PRODUCTS = 200
# Consume whole hyphenated codes before bare numeric IDs (never split a part number).
TOKEN = re.compile(r"(?<![\w.-])(?:[A-Za-z0-9]+(?:-[A-Za-z0-9]+)+|\d{5,12})(?![\w.-])")
HEADERS = {"productid", "productnumber", "itemnumber", "itemid", "partnumber", "sku", "fergusonid", "productcode"}


@dataclass
class Request:
    message: object
    ids: list[str]
    origins: dict[str, str]
    notes: list[tuple[str, str]]


def candidates(text):
    for line in text.splitlines():
        if re.match(r"\s*(?:From|To|Date|Sent|Subject|Message-ID):", line, re.I):
            continue
        line = re.sub(r"https?://\S+|\b\S+@\S+\b", "", line)
        for match in TOKEN.finditer(line):
            token = match.group().upper()
            if not any(c.isdigit() for c in token):
                continue
            if re.fullmatch(r"\d{4}-\d{2}-\d{2}", token):
                continue
            yield token


def attachment_ids(payload, filename):
    try:
        with ZipFile(BytesIO(payload)) as archive:
            if sum(x.file_size for x in archive.infolist()) > 30 * 1024 * 1024:
                raise ValueError(f"Attachment {filename} exceeds the expanded size limit")
        workbook = load_workbook(BytesIO(payload), read_only=True, data_only=False)
    except (BadZipFile, KeyError) as exc:
        raise ValueError(f"Cannot read Excel attachment {filename}") from exc
    try:
        for sheet in workbook:
            rows = sheet.iter_rows()
            columns = None
            for row_index, row in enumerate(rows, 1):
                if row_index > 10000:
                    raise ValueError(f"Attachment {filename} exceeds the row limit")
                if columns is None:
                    columns = [i for i, cell in enumerate(row)
                               if re.sub(r"[^a-z0-9]", "", str(cell.value).lower()) in HEADERS]
                    if not columns:
                        columns = None
                        if row_index >= 20:
                            raise ValueError(f"No product ID column found in {filename}/{sheet.title}")
                        continue
                    continue
                for column in columns:
                    if column >= len(row):
                        continue
                    cell = row[column]
                    if cell.data_type == "f":
                        raise ValueError(f"Formula in product ID cell {filename}/{sheet.title}!{cell.coordinate}")
                    if cell.value is None:
                        continue
                    value = str(cell.value).strip()
                    if isinstance(cell.value, float) and cell.value.is_integer():
                        value = str(int(cell.value))
                    # An explicit ID column also preserves short or unusual customer codes.
                    yield value.upper(), f"attachment:{filename}/{sheet.title}!{cell.coordinate}"
    finally:
        workbook.close()


def extract(raw: bytes) -> Request:
    if len(raw) > MAX_EMAIL_BYTES:
        raise ValueError("Email exceeds 15 MB")
    message = BytesParser(policy=policy.default).parsebytes(raw)
    result = Request(message, [], {}, [])

    def add(value, origin):
        if value in result.origins:
            result.notes.append((value, f"Duplicate removed: {origin}"))
            return
        if len(result.ids) >= MAX_PRODUCTS:
            raise ValueError(f"Request exceeds {MAX_PRODUCTS} products; split the request")
        result.ids.append(value)
        result.origins[value] = origin

    body = message.get_body(preferencelist=("plain", "html"))
    if body:
        content = body.get_content()
        if body.get_content_type() == "text/html":
            content = BeautifulSoup(content, "html.parser").get_text("\n")
        for value in candidates(content):
            add(value, "email:body")
    for part in message.walk():
        filename = part.get_filename()
        if not filename:
            continue
        if filename.lower().endswith(".xlsx"):
            for value, origin in attachment_ids(part.get_payload(decode=True), filename):
                add(value, origin)
        else:
            raise ValueError(f"Unsupported attachment {filename}; use XLSX for product lists")
    if not result.ids:
        raise ValueError("No product IDs found in the email body or Excel attachment")
    return result
