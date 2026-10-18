"""Build and validate the minimal Azure Functions deployment package."""

from __future__ import annotations

import argparse
import ast
import json
import sys
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent
DEFAULT_ZIP = ROOT / ".packages" / "func-deploy.zip"
REQUIRED_FILES = {
    "function_app.py",
    "host.json",
    "requirements.txt",
    "data/product_catalog.json",
    "data/product_catalog_sources.csv",
    "templates/rfq_template.xlsx",
}


def runtime_files() -> list[Path]:
    files = [ROOT / name for name in REQUIRED_FILES]
    files.extend(sorted((ROOT / "app").glob("*.py")))
    return sorted(files, key=lambda path: path.relative_to(ROOT).as_posix())


def validate_source(files: list[Path]) -> None:
    for path in files:
        if not path.is_file():
            raise FileNotFoundError(f"Required deployment file is missing: {path.relative_to(ROOT)}")
        if path.suffix == ".py":
            ast.parse(path.read_text(encoding="utf-8"), filename=str(path))
    json.loads((ROOT / "host.json").read_text(encoding="utf-8"))


def build_package(destination: Path) -> None:
    files = runtime_files()
    validate_source(files)
    destination.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(destination, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for path in files:
            archive.write(path, path.relative_to(ROOT).as_posix())
    validate_package(destination)
    print(f"Created {destination.relative_to(ROOT)} ({len(files)} files)")


def validate_package(path: Path) -> None:
    with zipfile.ZipFile(path) as archive:
        bad_file = archive.testzip()
        if bad_file:
            raise ValueError(f"ZIP integrity check failed at {bad_file}")
        names = set(archive.namelist())
        missing = REQUIRED_FILES - names
        if missing:
            raise ValueError(f"Deployment ZIP is missing required files: {', '.join(sorted(missing))}")
        forbidden = {".env", "local.settings.json"}
        if names & forbidden or any(name.startswith((".git/", ".azure/", ".tools/")) for name in names):
            raise ValueError("Deployment ZIP contains local credentials or deployment state")


def validate_workbook(path: Path) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=False)
    required_sheets = {"RFQ", "Sources", "Notes", "Summary"}
    missing = required_sheets - set(workbook.sheetnames)
    if missing:
        raise ValueError(f"Workbook is missing sheets: {', '.join(sorted(missing))}")
    workbook.close()
    print(f"Verified workbook: {path}")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=DEFAULT_ZIP)
    parser.add_argument("--verify-workbook", type=Path)
    args = parser.parse_args()
    try:
        if args.verify_workbook:
            validate_workbook(args.verify_workbook.resolve())
        else:
            build_package(args.output.resolve())
    except Exception as exc:
        print(f"Package check failed: {type(exc).__name__}: {exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
