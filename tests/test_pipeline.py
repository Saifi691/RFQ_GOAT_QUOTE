import json
from email import policy
from email.message import EmailMessage
from email.parser import BytesParser
from io import BytesIO
from pathlib import Path

import azure.functions as func
import pytest
from openpyxl import Workbook, load_workbook

from app.config import ROOT, Settings
from app.extract import extract
from app.lookup import Lookup, parse_pdf_text, parse_product
from app.models import Evidence, Product
from app.pipeline import build
from app.quality import check_products
from app.workbook import validate_workbook

EXPECTED = {
    1: ['5106912', '5106935', '5106914', '5106966', '5106891', '5106880'],
    2: ['1376044', '1432040', '233900', '4742155', '5106899', '9677514', '7099537', '225123'],
    3: ['5106878', '5106907', '4570880', '9677531', '17783-79-3000', '31-024', 'MTX-4409-XQ'],
    4: ['5106926', '8067331', '1844552', '4783795', '5106888', '9677533', '17774-79-3000'],
}


def sample(number):
    return next((ROOT / 'emails').glob(f'email-{number}-*.eml')).read_bytes()


@pytest.mark.parametrize('number', [1, 2, 3, 4])
def test_sample_ids(number):
    assert extract(sample(number)).ids == EXPECTED[number]


def test_reference_exact_rfq_and_reply(tmp_path):
    result = build(sample(1), tmp_path, Settings(offline=True))
    actual = load_workbook(result['workbook'])
    expected = load_workbook(ROOT / 'reference/email-1 reference output.xlsx')
    assert list(actual['RFQ'].values) == list(expected['RFQ'].values)
    assert {'RFQ', 'Sources', 'Notes', 'Summary', 'Commercial'} <= set(actual.sheetnames)
    validate_workbook(result['workbook'], EXPECTED[1])
    reply = BytesParser(policy=policy.default).parsebytes(result['reply'].read_bytes())
    assert reply['In-Reply-To'] == '<rfq-sample-1@sourcing.example.com>'
    assert reply['Auto-Submitted'] == 'auto-replied'
    attachment = next(reply.iter_attachments())
    assert attachment.get_filename() == 'RFQ.xlsx'
    assert attachment.get_payload(decode=True) == result['workbook'].read_bytes()
    actual.close()
    expected.close()


@pytest.mark.parametrize('number, missing', [(2, ['225123']),
    (3, ['17783-79-3000', '31-024', 'MTX-4409-XQ']), (4, ['9677533', '17774-79-3000'])])
def test_unknowns_retained_and_blank(tmp_path, number, missing):
    result = build(sample(number), tmp_path, Settings(offline=True))
    workbook = load_workbook(result['workbook'])
    rows = list(workbook['RFQ'].values)[1:]
    assert [r[0] for r in rows] == EXPECTED[number]
    assert [r[0] for r in rows if r[-1] == 'Not Found'] == missing
    for row in rows:
        if row[-1] == 'Not Found':
            assert all(value is None for value in row[1:-1])
    workbook.close()


def test_html_and_duplicate():
    msg = EmailMessage()
    msg.set_content('<p>Item #5106912</p><p>Repeat 5106912; MTX-4409-XQ</p>', subtype='html')
    request = extract(msg.as_bytes())
    assert request.ids == ['5106912', 'MTX-4409-XQ']
    assert len(request.notes) == 1


def test_attachment_formula_rejected():
    book = Workbook()
    book.active.append(['Item number'])
    book.active.append(['=12345'])
    data = BytesIO()
    book.save(data)
    msg = EmailMessage()
    msg.set_content('See attached list')
    msg.add_attachment(data.getvalue(), maintype='application', subtype='octet-stream', filename='items.xlsx')
    with pytest.raises(ValueError, match='Formula'):
        extract(msg.as_bytes())


def test_website_wins_and_catalog_fills_only_blank_fields(monkeypatch):
    lookup = Lookup(Settings(pdfs=False))
    monkeypatch.setattr(lookup, 'live', lambda *_: {
        'Description': Evidence('Live description', 'ferguson_page', 'https://www.ferguson.com/'),
        'Length': Evidence('1000 ft', 'ferguson_page', 'https://www.ferguson.com/')})
    result = lookup.lookup('5106912')
    assert result.found
    assert result.fields['Description'].value == 'Live description'
    assert result.fields['Length'].source_type == 'ferguson_page'
    assert result.fields['ASTM'].source_type == 'catalog'
    assert 'original spec_sheet' in result.fields['ASTM'].source
    # Website beats catalog; a real difference is recorded, a formatting-only one is not.
    assert any("Conflict for Description" in n for n in result.notes)
    assert not any("Conflict for Length" in n for n in result.notes)
    assert any('filled' in n and 'showed them' in n for n in result.notes)
    lookup.close()


def test_product_removed_from_website_is_not_revived_from_catalog(monkeypatch):
    from app.browser import NotOnSite
    lookup = Lookup(Settings(pdfs=False))
    def gone(*_):
        raise NotOnSite('ferguson.com has no product page for this ID')
    monkeypatch.setattr(lookup, 'live', gone)
    result = lookup.lookup('5106912')
    assert not result.found and result.fields == {}
    lookup.close()


def test_live_failure_uses_catalog(monkeypatch):
    lookup = Lookup(Settings(pdfs=False))
    def fail(*args):
        raise TimeoutError('timeout')
    monkeypatch.setattr(lookup, 'live', fail)
    result = lookup.lookup('5106912')
    assert result.found
    assert result.fields['Length'].source_type == 'catalog'
    assert any('historical' in note for note in result.notes)
    lookup.close()


def test_priority_and_conflict():
    product = Product('1')
    product.merge({'Material': Evidence('Plastic', 'ferguson_page', 'page')})
    product.merge({'Material': Evidence('Steel', 'spec_sheet', 'pdf')})
    assert product.fields['Material'].value == 'Plastic'
    assert 'Conflict' in product.notes[0]


def test_live_parser_requires_exact_identity():
    schema = {'@type': 'Product', 'productID': '5106912', 'name': 'Tape',
        'offers': {'price': '10', 'priceCurrency': 'USD', 'availability': 'https://schema.org/InStock'},
        'additionalProperty': [{'name': 'Length', 'value': '1000 ft'}]}
    html = '<script type="application/ld+json">' + json.dumps(schema) + '</script>'
    result = parse_product(html, 'https://www.ferguson.com/product/5106912.html', '5106912')
    assert result['Length'].value == '1000 ft'
    assert result['Price (USD)'].value == '10'
    with pytest.raises(ValueError, match='exact product'):
        parse_product(html, 'https://www.ferguson.com/product/999999.html', '999999')


def test_pdf_exact_model_is_high_confidence():
    fields, note = parse_pdf_text('Model SD6105B52\nThickness: 5 mil\nMaterial: Polyethylene Film',
                                  '5106912', 'SD6105B52', 'pdf')
    assert fields['Thickness'].value == '5 mil' and fields['Thickness'].confidence == 'high'
    assert fields['Material Type'].value == 'Polyethylene'
    assert note == ''


def test_pdf_family_sheet_is_medium_and_flagged():
    # Layout of the real Presco/Trident SD family sheet: value on the line after the label.
    text = 'Base Sku: SD\nStandard Widths:\n2” | 3” | 4” | 6” | 12”\nStandard Length:\n1,000’\nOverall Thickness:\n5 mil'
    fields, note = parse_pdf_text(text, '5106912', 'SD6105B52', 'pdf')
    assert {k: v.value for k, v in fields.items()} == {'Length': '1000 ft', 'Thickness': '5 mil'}
    assert all(v.confidence == 'medium' for v in fields.values())
    assert 'Review spec sheet values' in note
    fields, note = parse_pdf_text(text, '5106912', 'B3104Y3', 'pdf')
    assert fields == {} and 'not used' in note


def test_pdf_skips_option_lists_and_conflicts_and_reads_standards():
    text = ('B3104Y3\nGauge(s): 2, 3, 4, 6, & 7 mil (+/- 10%)\nWidth: 3 in\nWidth: 6 in\n'
            'Occupational Safety and Health Administration (OSHA): Regulation 1910-144(a)\n'
            'Occupational Safety and Health Administration (OSHA): Regulation 1926-956(c)(1)\n'
            'Meets ASTM D-2103 and ASTM D882')
    fields, _ = parse_pdf_text(text, '5106880', 'B3104Y3', 'pdf')
    assert 'Thickness' not in fields and 'Width' not in fields
    assert fields['OSHA'].value == 'OSHA 1910-144(a);OSHA 1926-956(c)(1)'
    assert fields['ASTM'].value == 'ASTM D-2103|ASTM D-882'


FERGUSON_PAGE = '''<html><head><title>Presco Tape | Ferguson</title>
<script type="application/ld+json">{"@type": "Product", "url": "https://www.ferguson.com/product/x/5106912.html",
 "name": "6 in x 1000 ft Blue Detector Tape", "sku": "PSD6105B52", "brand": {"name": "Presco"},
 "image": "https://i.ferg-img.com/i/fergusonprod/5106912_1815492_primary",
 "description": ["Durable construction", "Permanently imprinted"],
 "offers": {"price": "139.02", "priceCurrency": "USD", "availability": "InStock"}}</script></head><body>
<div>Part #PSD6105B52</div><div>Item #5106912</div><div>Manufacturer Part #SD6105B52</div>
<a href="/category/tools/diagnostic%2C-inspection-locating-tools/">Diagnostic, Inspection</a>
<div>Documents</div><a href="https://fergusonprod.a.bigcontent.io/v1/static/5106912_17560988_specification">SPECIFICATION</a>
<div class="product-specification-table">
 <div class="d-flex c-product-details__specs__specsAttributes"><div class="font-weight-bold">Thickness</div><div> 5 mil </div></div>
 <div class="d-flex c-product-details__specs__specsAttributes"><div class="font-weight-bold">Color Finish Name</div><div>Blue</div></div>
 <div class="d-flex c-product-details__specs__specsAttributes"><div class="font-weight-bold">Material</div><div>Aluminum Foil, Plastic</div></div>
 <div class="d-flex c-product-details__specs__specsAttributes"><div class="font-weight-bold">Application</div><div>Water, Marking, Underground</div></div>
</div></body></html>'''


def test_live_parser_reads_ferguson_spec_block():
    url = 'https://www.ferguson.com/product/x/5106912.html'
    fields = {k: v.value for k, v in parse_product(FERGUSON_PAGE, url, '5106912').items()}
    assert fields['Thickness'] == '5 mil'
    assert fields['Color Finish'] == 'Blue'
    assert fields['Material'] == 'Aluminum Foil|Plastic'
    assert fields['Application'] == 'Water|Marking|Underground'
    assert fields['Product details'] == 'Durable construction|Permanently imprinted'
    assert fields['Manufacturer Part #'] == 'SD6105B52'
    assert fields['Ferguson Part #'] == 'PSD6105B52'
    # "Inspection" contains "spec"; only the real specification document may be chosen.
    assert fields['Spec Sheet URL'] == 'https://fergusonprod.a.bigcontent.io/v1/static/5106912_17560988_specification'


def test_live_parser_page_without_json_ld_needs_visible_item_number():
    page = ('<h1>1000 ft. Detectable Gas Tape</h1><div>Part #PDTGU</div><div>Item #233900</div>'
            '<div class="c-product-details__specs__specsAttributes"><div>Width</div><div>6 in</div></div>')
    url = 'https://www.ferguson.com/product/pro-pak-industries-1000-ft.-detectable-gas-tape-pdtgu/233900.html'
    fields = {k: v.value for k, v in parse_product(page, url, '233900').items()}
    assert fields['Description'] == '1000 ft. Detectable Gas Tape' and fields['Width'] == '6 in'
    with pytest.raises(ValueError, match='exact product'):
        parse_product(page.replace('233900', '233901'), url, '233900')


def test_live_mode_fails_fast_without_chrome(tmp_path):
    from app.browser import BrowserUnavailable
    settings = Settings(mode='live', cdp_url='http://127.0.0.1:9', pdfs=False)
    with pytest.raises(BrowserUnavailable, match='remote-debugging-port=9333'):
        build(sample(1), tmp_path, settings)


def test_background_mode_without_chrome_falls_back(tmp_path, monkeypatch):
    import app.browser as browser
    monkeypatch.setattr(browser, 'find_chrome', lambda *_: None)
    monkeypatch.setattr(browser, 'cdp_ready', lambda *_args, **_kw: False)
    import app.lookup as lookup
    monkeypatch.setattr(lookup.Lookup, 'live', lambda *_: (_ for _ in ()).throw(TimeoutError('no network in test')))
    result = build(sample(1), tmp_path, Settings(mode='background', pdfs=False))
    notes = list(load_workbook(result['workbook'])['Notes'].values)
    assert any('Browser unavailable' in str(row[1]) for row in notes)
    assert result['summary']['Products Found'] == 6


def test_units_keep_evidence_and_do_not_guess():
    product = Product('1', {'Length': Evidence('1000 Feet', 'spec_sheet', 'pdf'),
                            'Width': Evidence('3', 'ferguson_page', 'page')})
    check_products([product])
    assert product.fields['Length'].value == '1000 ft'
    assert 'normalized from' in product.fields['Length'].source
    assert product.fields['Width'].value == '3'


def test_ferguson_quote_markup_and_number_formats():
    from app.lookup import page_text, same_value
    assert page_text("Pro-Line<CQUOTA/>s tape") == "Pro-Line's tape"
    assert same_value('Price (USD)', '112.40', '112.4')
    assert same_value('Length', '1000 ft', '1000')
    assert not same_value('Lettering', 'Caution Buried Reclaimed Water Line Below', 'Caution Buried Water Line Below')


def test_azure_and_local_match(tmp_path):
    from function_app import rfq, healthz
    request = func.HttpRequest(method='POST', url='http://localhost/rfq', body=sample(1), params={'offline': 'true'})
    response = rfq(request)
    assert response.status_code == 200
    cloud = load_workbook(BytesIO(response.get_body()))
    local_result = build(sample(1), tmp_path, Settings(offline=True))
    local = load_workbook(local_result['workbook'])
    for name in local.sheetnames:
        assert list(local[name].values) == list(cloud[name].values)
    assert healthz(request).status_code == 200
    cloud.close()
    local.close()


def test_offline_never_uses_network(tmp_path, monkeypatch):
    import requests
    def fail(*args, **kwargs):
        pytest.fail('Offline build attempted network access')
    monkeypatch.setattr(requests.Session, 'get', fail)
    build(sample(1), tmp_path, Settings(offline=True))
