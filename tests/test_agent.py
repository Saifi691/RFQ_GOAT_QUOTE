from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from app.agent import QualityAgent, supported
from app.config import ROOT, Settings
from app.models import Conflict, Evidence, Product
from app.pipeline import build
from app.review import decide, load_state


def sample(number):
    return next((ROOT / 'emails').glob(f'email-{number}-*.eml')).read_bytes()


def page(value):
    return Evidence(value, 'ferguson_page', 'https://www.ferguson.com/product/x.html')


def sewer_tape():
    # Real case: ferguson.com lists 1432040 "6 in. x 300 ft." with Length "300 in".
    return Product('1432040', {
        'Description': page('6 in. x 300 ft. Detectable Sewer Tape'),
        'Length': page('300 in'), 'Width': page('6 in'), 'Color Finish': page('Green'),
        'Lettering': page('Caution Buried Sewer Line Below'), 'Application': page('Gas|Marking')},
        found=True, snapshots={'ferguson_page': 'Title: 6 in. x 300 ft. Detectable Sewer Tape\nLength: 300 in'})


def test_rules_plan_flags_description_mismatch_with_evidence_backed_fix():
    product = sewer_tape()
    agent = QualityAgent([product])
    agent.run_rules()
    fix = next(f for f in agent.findings if f.field == 'Length')
    assert fix.severity == 'critical' and fix.proposed_value == '300 ft'
    assert '300 ft' in fix.evidence
    assert product.fields['Length'].value == '300 in'  # the agent never writes values itself
    assert not any(f.field == 'Width' and f.verdict == 'issue' for f in agent.findings)


def test_apwa_colour_code():
    product = sewer_tape()
    product.fields['Lettering'] = page('Caution Buried Water Line Below')
    problems = QualityAgent([product]).check_standard('1432040')['problems']
    messages = ' '.join(p['message'] for p in problems)
    assert 'potable water (APWA blue)' in messages      # lettering says water, tape is green
    assert 'Gas (APWA yellow)' in messages              # gas application on green tape
    product.fields['Lettering'] = page('Caution Buried Sewer Line Below')
    product.fields['Application'] = page('Water|Marking|Underground')
    assert QualityAgent([product]).check_standard('1432040')['problems'] == []


def test_verdicts_follow_how_directly_the_source_supports_the_value():
    product = Product('5106912', {
        'Length': page('1000 ft'),
        'ASTM': Evidence('ASTM D-882', 'spec_sheet', 'pdf (family)', 'medium'),
        'UL Listed': Evidence('No', 'catalog', 'catalog; original validated_rule: every tape says No'),
        'Weight': Evidence('14.25 lb', 'catalog', 'catalog; original ferguson_rfq_sample: sample RFQ')}, found=True)
    verdicts = QualityAgent([product]).verify_sources('5106912')
    assert {k: v['verdict'] for k, v in verdicts.items()} == {
        'Length': 'verified', 'ASTM': 'inferred', 'UL Listed': 'inferred', 'Weight': 'verified'}


def test_fix_must_come_from_a_source():
    agent = QualityAgent([sewer_tape()])
    agent.inspect_product('1432040')
    with pytest.raises(ValueError, match='never guess'):
        agent.propose_fix('1432040', 'Thickness', '5 mil', 'my opinion')
    assert agent.propose_fix('1432040', 'Length', '300 ft', 'Title says 300 ft', reason='Unit is wrong')['accepted']
    assert agent.findings[-1].message == 'Unit is wrong'
    assert not agent.propose_fix('1432040', 'Length', '300 ft', 'again')['accepted']  # one proposal per fix
    assert supported('1000 ft', 'Standard Length: 1,000’') and not supported('900 ft', 'Length: 1,000’')


def test_conflicts_become_findings():
    product = Product('9677531', {'Lettering': page('Caution Buried Reclaimed Water Line Below')}, found=True)
    product.conflicts.append(Conflict('Lettering', product.fields['Lettering'],
                                      Evidence('Caution Buried Water Line Below', 'catalog', 'catalog')))
    agent = QualityAgent([product])
    agent.run_rules()
    conflict = next(f for f in agent.findings if f.verdict == 'conflict')
    assert conflict.severity == 'warning' and 'Reclaimed' in conflict.message


def test_ai_planner_uses_tools_and_cannot_write_values():
    product = sewer_tape()
    calls = [
        [SimpleNamespace(type='function_call', name='inspect_product', arguments='{"product_id":"1432040"}', call_id='1'),
         SimpleNamespace(type='function_call', name='read_source',
                         arguments='{"product_id":"1432040","source":"ferguson_page"}', call_id='2')],
        [SimpleNamespace(type='function_call', name='propose_fix', call_id='3',
                         arguments='{"product_id":"1432040","field":"Width","value":"9 in","evidence":"guess"}'),
         SimpleNamespace(type='function_call', name='record_finding', call_id='4',
                         arguments='{"product_id":"1432040","field":"Application","severity":"warning",'
                                   '"message":"Gas on sewer tape","evidence":"Application: Gas|Marking"}')],
    ]
    responses = iter([SimpleNamespace(output=c, output_text='') for c in calls] +
                     [SimpleNamespace(output=[], output_text='Two problems found.')])
    client = SimpleNamespace(responses=SimpleNamespace(create=lambda **kwargs: next(responses)))
    agent = QualityAgent([product])
    assert agent.run_ai(client=client, model='test-model') == 'Two problems found.'
    tools = [step['tool'] for step in agent.trace]
    assert tools == ['inspect_product', 'read_source', 'propose_fix', 'record_finding']
    assert 'never guess' in agent.trace[2]['result']          # unsupported fix rejected
    assert any(f.found_by == 'ai' and f.field == 'Application' for f in agent.findings)
    assert product.fields['Width'].value == '6 in'


def test_reviewer_approval_regenerates_workbook(tmp_path, monkeypatch):
    # Email 1 offline is clean; inject the real 1432040 mismatch into its first product.
    import app.pipeline as pipeline
    original = pipeline.check_products
    def with_mismatch(products):
        products[0].fields['Length'] = page('1000 in')
        return original(products)
    monkeypatch.setattr(pipeline, 'check_products', with_mismatch)
    result = build(sample(1), tmp_path, Settings(offline=True))
    number = next(i for i, f in enumerate(result['findings'], 1) if f.proposed_value)
    assert result['findings'][number - 1].proposed_value == '1000 ft'

    finding, summary = decide(tmp_path, number, 'approved')
    assert finding.decision == 'approved'
    workbook = load_workbook(tmp_path / 'RFQ.xlsx')
    header = [c.value for c in workbook['RFQ'][1]]
    assert workbook['RFQ'].cell(2, header.index('Length') + 1).value == '1000 ft'
    sources = [r for r in workbook['Sources'].values if r[0] == '5106912' and r[1] == 'Length']
    assert sources[0][3] == 'agent_fix' and 'Reviewer approved' in sources[0][5]
    review_rows = list(workbook['Agent Review'].values)
    assert any(r[0] == number and r[-1] == 'approved' for r in review_rows)
    assert (tmp_path / 'RFQ-original.xlsx').exists()
    assert summary['Agent: fixes approved'] == 1
    with pytest.raises(ValueError, match='Already'):
        decide(tmp_path, number, 'rejected')
    _, _, products, _ = load_state(tmp_path)
    assert products[0].fields['Length'].source_type == 'agent_fix'


def test_email_1_reference_is_clean_and_rfq_sheet_unchanged(tmp_path):
    result = build(sample(1), tmp_path, Settings(offline=True))
    summary = result['agent_summary']
    assert summary['issues'] == 0 and summary['critical'] == 0
    assert summary['verified'] > summary['inferred'] > 0
    expected = load_workbook(ROOT / 'reference/email-1 reference output.xlsx')
    actual = load_workbook(result['workbook'])
    assert list(actual['RFQ'].values) == list(expected['RFQ'].values)
    assert 'Agent Review' in actual.sheetnames
